"""The portal deck: problem, what it does, capabilities, and an annotated tour (shots/ + shots/boxes.json)."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

D = Path(__file__).resolve().parent
SHOTS = D / "shots"
BOXES = json.loads((SHOTS / "boxes.json").read_text())
VER = (D.parents[1] / "VERSION").read_text().strip()

NAVY = RGBColor(0x0B, 0x1B, 0x2E)
TEAL = RGBColor(0x0E, 0x74, 0x90)
TEAL_L = RGBColor(0xE0, 0xF2, 0xF6)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x67, 0x78)
LINE = RGBColor(0xD6, 0xDD, 0xE5)
BG = RGBColor(0xF6, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x15, 0x80, 0x3D)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def box(sl, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def text(sl, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: a string, or a list of paragraphs; a paragraph is a string or a list of (text, {size, color, bold})."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            f = r.font; f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color); f.italic = o.get("italic", False)
    return tb


def shape_text(s, t, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER):
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    lines = t.split("\n")
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = l
        r.font.name = FONT; r.font.size = Pt(size if i == 0 else size - 2); r.font.bold = bold if i == 0 else False
        r.font.color.rgb = color


def header(sl, title, kicker=None, n=None):
    box(sl, 0, 0, SW, 7.5, fill=WHITE)
    box(sl, 0, 0, 0.14, 7.5, fill=TEAL)
    if kicker:
        text(sl, 0.45, 0.28, 9, 0.3, kicker.upper(), size=11, color=TEAL, bold=True)
    text(sl, 0.45, 0.5, 12.2, 0.6, title, size=26, color=NAVY, bold=True)
    if n:
        text(sl, 12.2, 7.08, 0.9, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    text(sl, 0.45, 7.08, 6, 0.3, f"C8000v DMVPN Portal · v{VER}", size=10, color=MUTED)


def notes(sl, t):
    sl.notes_slide.notes_text_frame.text = t


# ---------- 1. title ----------
def title_slide():
    sl = prs.slides.add_slide(BLANK)
    box(sl, 0, 0, SW, SH, fill=NAVY)
    box(sl, 0, 5.2, SW, 0.06, fill=TEAL)
    # a small DMVPN motif: three hubs, spokes, one shortcut
    hubs = [(9.3, 1.3), (10.8, 1.0), (12.3, 1.3)]
    spokes = [(8.9, 3.6), (10.0, 4.1), (11.2, 4.2), (12.4, 3.7)]
    for hx, hy in hubs:
        for sx, sy in spokes:
            c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(hx), Inches(hy), Inches(sx), Inches(sy))
            c.line.color.rgb = RGBColor(0x2B, 0x4C, 0x6B); c.line.width = Pt(1)
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(10.0), Inches(4.1), Inches(12.4), Inches(3.7))
    c.line.color.rgb = RGBColor(0xA7, 0x8B, 0xFA); c.line.width = Pt(3)
    for hx, hy in hubs:
        s = box(sl, hx - 0.22, hy - 0.22, 0.44, 0.44, fill=TEAL, shape=MSO_SHAPE.OVAL)
    for sx, sy in spokes:
        box(sl, sx - 0.16, sy - 0.16, 0.32, 0.32, fill=RGBColor(0x94, 0xA3, 0xB8), shape=MSO_SHAPE.OVAL)
    text(sl, 0.8, 1.6, 8, 0.4, "NETWORK AS CODE · DMVPN PHASE 3", size=13, color=RGBColor(0x5E, 0xC8, 0xDB), bold=True)
    text(sl, 0.8, 2.05, 8.2, 1.4, "C8000v DMVPN Portal", size=46, color=WHITE, bold=True)
    text(sl, 0.8, 3.05, 7.6, 1.5, "One place to see, change, prove and share a multi-hub DMVPN service: "
         "for the operators who run it and the customers who use it.", size=18, color=RGBColor(0xCB, 0xD5, 0xE1))
    text(sl, 0.8, 5.55, 11, 0.4, "Three Catalyst 8000v hubs · C8000v and VyOS customers · two providers · Nautobot · Terraform (NAC) · VictoriaMetrics",
         size=13, color=RGBColor(0x94, 0xA3, 0xB8))
    text(sl, 0.8, 6.0, 8, 0.4, f"Version {VER} · September 2026", size=13, color=RGBColor(0x94, 0xA3, 0xB8))
    notes(sl, "Title. The portal runs on the lab host at :8094 and fronts the whole lab: lab.conf is the source of truth, "
              "Network-as-Code pushes the Catalyst 8000v routers, SSH pushes the VyOS routers, Nautobot mirrors the model.")


# ---------- 2. problem ----------
def problem_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "Running a DMVPN service by hand does not scale", "The problem", n)
    text(sl, 0.45, 1.2, 12.3, 0.6, "A single customer touches three hubs, one or two providers, two router operating systems and a source of truth. "
         "Every question and every change means logging in to several boxes.", size=15, color=MUTED)
    items = [
        ("Scattered state", "Is every spoke registered with every hub? Are shortcuts forming? Health lives in a dozen CLI sessions on IOS XE and VyOS."),
        ("Slow, risky changes", "Adding, changing or removing a customer means consistent edits on hubs, spoke, provider and Nautobot. It is easy to miss one."),
        ("Drift", "Hand edits leave routers differing from the intended design, and nobody notices until something breaks."),
        ("No change control", "Who changed what, who approved it, and what exactly changed on each router? Usually nobody can say."),
        ("Unproven resilience", "Dual hubs and dual providers look good on paper, but how long does failover really take?"),
        ("Blind customers", "Customers cannot see their own service, test it, or ask for changes. SLA reports are assembled by hand."),
        ("Static secrets", "The DMVPN pre-shared key is set once and rarely rotated, because rotating it by hand risks an outage."),
    ]
    x0, y0, cw, ch, gap = 0.45, 2.05, 3.0, 2.2, 0.12
    for i, (h, b) in enumerate(items):
        col, row = i % 4, i // 4
        x, y = x0 + col * (cw + gap), y0 + row * (ch + gap)
        box(sl, x, y, cw, ch, fill=BG, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        box(sl, x, y + 0.18, 0.07, 0.42, fill=ORANGE)
        text(sl, x + 0.2, y + 0.16, cw - 0.3, 0.45, h, size=15, color=NAVY, bold=True)
        text(sl, x + 0.2, y + 0.62, cw - 0.32, ch - 0.7, b, size=12, color=INK)
    x, y = x0 + 3 * (cw + gap), y0 + ch + gap
    s = box(sl, x, y, cw, ch, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    text(sl, x + 0.2, y + 0.2, cw - 0.4, ch - 0.4, [[("The result", {"bold": True, "size": 15, "color": RGBColor(0x5E, 0xC8, 0xDB)})],
         [("Slow onboarding, outages from avoidable mistakes, and no evidence for customers or auditors.", {"size": 13, "color": WHITE})]], spacing=6)
    notes(sl, "The problem statement: the service is simple to draw but has many moving parts, and without tooling every question costs a round of logins.")


# ---------- 3. what it is ----------
def what_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "What the portal does", "The solution", n)
    text(sl, 0.45, 1.2, 12.3, 0.7, [[("One web portal on top of a single source of truth (", {}), ("lab.conf", {"bold": True}),
         ("). It reads every router live, turns every change into a controlled, verified job, and gives each customer a view of its own service.", {})]],
         size=15, color=MUTED)

    def node(x, y, w, h, title, sub, fill=WHITE, line=TEAL, tc=NAVY):
        s = box(sl, x, y, w, h, fill=fill, line=line, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(sl, x + 0.08, y + 0.1, w - 0.16, 0.35, title, size=13.5, color=tc, bold=True, align=PP_ALIGN.CENTER)
        text(sl, x + 0.08, y + 0.44, w - 0.16, h - 0.5, sub, size=10.5, color=MUTED if tc == NAVY else RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.CENTER)
        return s

    def arrow(x1, y1, x2, y2, label=None, lx=None, ly=None):
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = MUTED; c.line.width = Pt(1.5)
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
        if label:
            text(sl, lx if lx is not None else (x1 + x2) / 2 + 0.05, ly if ly is not None else (y1 + y2) / 2 - 0.15, 2.2, 0.3, label, size=9.5, color=MUTED)

    node(0.5, 2.3, 2.4, 1.05, "Operators", "viewer · operator · approver · admin")
    node(0.5, 4.3, 2.4, 1.05, "Customers", "their own site, SLA, requests, tests")
    node(3.8, 2.6, 2.9, 2.5, "Portal  :8094", "FastAPI · jobs · change control · SLA probes · failover lab · config history · roles",
         fill=NAVY, line=NAVY, tc=WHITE)
    arrow(2.9, 2.83, 3.8, 3.3); arrow(2.9, 4.83, 3.8, 4.4)
    node(7.6, 1.95, 2.35, 1.0, "lab.conf", "the intended design")
    node(7.6, 3.35, 2.35, 1.0, "Render", "hub, spoke and provider config")
    node(7.6, 4.75, 2.35, 1.0, "Nautobot", "mirrors the model")
    arrow(6.7, 3.2, 7.6, 2.45); arrow(8.775, 2.95, 8.775, 3.35); arrow(6.7, 4.3, 7.6, 5.2)
    node(10.6, 1.55, 2.3, 1.05, "Catalyst 8000v", "Terraform (NAC) → RESTCONF")
    node(10.6, 2.95, 2.3, 1.05, "VyOS", "SSH push, declarative")
    node(10.6, 4.35, 2.3, 1.05, "Providers", "mpls and mpls2 (VyOS)")
    node(10.6, 5.75, 2.3, 0.9, "VictoriaMetrics", "SLA and application probes")
    arrow(9.95, 3.75, 10.6, 2.1); arrow(9.95, 3.85, 10.6, 3.45); arrow(9.95, 3.95, 10.6, 4.85)
    arrow(6.7, 4.9, 10.6, 6.2, "every minute: pings, app checks", 7.4, 6.2)
    text(sl, 0.5, 5.95, 6.3, 0.9, [[("Read path: ", {"bold": True, "color": TEAL}), ("the portal polls every router (RESTCONF, CLI, VyOS op mode) and draws what is really there, not what should be.", {})]],
         size=11.5, color=INK)
    notes(sl, "Architecture: the portal never edits routers ad hoc. Changes go through lab.conf, the renderer and the same push pipeline, then a verification step.")


# ---------- 4. capabilities ----------
def caps_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "Capabilities at a glance", "What you can do", n)
    caps = [
        ("See", "Live cloud table and KPIs, a topology map drawn from live state, per-customer details, path trace and 'watch the shortcut form'."),
        ("Change", "Add, modify and remove customers, deploy the model, fix drift, back up and restore, choose the preferred hub or add a second provider."),
        ("Control", "Change requests with four-eyes approval, change windows, emergency changes, and maintenance mode that mutes alerts."),
        ("Audit", "Every job snapshots every router before and after it runs, records who asked and who approved, and shows a diff per router."),
        ("Prove", "SLA per customer, application checks, failure simulation with measured failover times, and capacity: hub load and room to grow."),
        ("Secure", "Sign-in with roles, admin-only pre-shared key rotation with a random key, and secrets masked in every view and log."),
        ("Customers", "A read-only customer portal: its own site and applications, SLA and PDF report, change requests and self-service tests."),
        ("Automate", "A REST API for everything (Swagger), Prometheus metrics, progress and ETA for each job, and a 73-test Robot suite after each change."),
    ]
    x0, y0, cw, ch, gx, gy = 0.45, 1.3, 3.0, 2.72, 0.12, 0.14
    for i, (h, b) in enumerate(caps):
        col, row = i % 4, i // 4
        x, y = x0 + col * (cw + gx), y0 + row * (ch + gy)
        box(sl, x, y, cw, ch, fill=BG, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        c = box(sl, x + 0.22, y + 0.24, 0.5, 0.5, fill=TEAL, shape=MSO_SHAPE.OVAL)
        shape_text(c, str(i + 1), size=15)
        text(sl, x + 0.86, y + 0.3, cw - 1, 0.4, h, size=18, color=NAVY, bold=True)
        text(sl, x + 0.22, y + 0.92, cw - 0.4, ch - 1.0, b, size=12.5, color=INK)
    notes(sl, "Eight capability areas. The tour that follows shows each one in the portal, with numbered callouts matching the explanation beside each screenshot.")


def divider(title, sub, n):
    sl = prs.slides.add_slide(BLANK)
    box(sl, 0, 0, SW, SH, fill=NAVY)
    box(sl, 0.8, 3.05, 0.12, 1.3, fill=ORANGE)
    text(sl, 1.15, 2.95, 11, 0.8, title, size=36, color=WHITE, bold=True)
    text(sl, 1.15, 3.75, 11, 0.8, sub, size=17, color=RGBColor(0xCB, 0xD5, 0xE1))
    text(sl, 12.2, 7.08, 0.9, 0.3, str(n), size=10, color=RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.RIGHT)
    return sl


# ---------- annotated screenshots ----------
IMG_X, IMG_Y, IMG_W = 0.45, 1.3, 8.75
IMG_H = IMG_W * 1000 / 1600
PX = 9.5


def shot_slide(n, name, kicker, title, points, takeaway, extra_boxes=None, drop=()):
    """points: explanations, one per callout, in callout order (after extra boxes are merged)."""
    sl = prs.slides.add_slide(BLANK)
    header(sl, title, kicker, n)
    box(sl, IMG_X + 0.04, IMG_Y + 0.06, IMG_W, IMG_H, fill=LINE)                       # a soft shadow
    pic = sl.shapes.add_picture(str(SHOTS / f"{name}.png"), Inches(IMG_X), Inches(IMG_Y), Inches(IMG_W), Inches(IMG_H))
    pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
    calls = [c for i, c in enumerate(BOXES.get(name, [])) if i not in drop] + (extra_boxes or [])
    assert len(calls) == len(points), f"{name}: {len(calls)} callouts, {len(points)} explanations"
    for i, c in enumerate(calls):
        bx, by, bw, bh = c["box"]
        pad = 0.004
        x, y = IMG_X + (bx - pad) * IMG_W, IMG_Y + (by - pad) * IMG_H
        w, h = (bw + 2 * pad) * IMG_W, max((bh + 2 * pad) * IMG_H, 0.16)
        x, y = max(IMG_X, x), max(IMG_Y, y)
        w, h = min(w, IMG_X + IMG_W - x), min(h, IMG_Y + IMG_H - y)
        r = box(sl, x, y, w, h, line=ORANGE, lw=2.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08 if h > 0.5 else 0.25)
        d = 0.34
        cx = min(max(x - d / 2, IMG_X - d / 2 + 0.02), IMG_X + IMG_W - d)
        cy = min(max(y - d / 2, IMG_Y - d / 2 + 0.02), IMG_Y + IMG_H - d)
        b = box(sl, cx, cy, d, d, fill=ORANGE, line=WHITE, lw=1.5, shape=MSO_SHAPE.OVAL)
        shape_text(b, str(i + 1), size=12)
    # explanation panel: the points, then the takeaway pinned to the bottom of the screenshot
    ty = IMG_Y + IMG_H - 1.0
    hs = [0.26 + (max(1, -(-len(p[1]) // 40)) * 0.2 if isinstance(p, tuple) and p[1] else 0) for p in points]
    gap = max(0.08, min(0.3, (ty - 0.1 - IMG_Y - sum(hs)) / max(1, len(points) - 1)))
    py = IMG_Y
    for i, p in enumerate(points):
        head, body = p if isinstance(p, tuple) else (p, "")
        b = box(sl, PX, py + 0.02, 0.3, 0.3, fill=ORANGE, shape=MSO_SHAPE.OVAL)
        shape_text(b, str(i + 1), size=11)
        text(sl, PX + 0.42, py - 0.02, SW - PX - 0.75, hs[i] + 0.05, [[(head, {"bold": True, "color": NAVY, "size": 13})], [(body, {"size": 11.5})]], size=11.5)
        py += hs[i] + gap
    box(sl, PX, ty, SW - PX - 0.4, IMG_Y + IMG_H - ty, fill=TEAL_L, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(sl, PX + 0.15, ty + 0.08, SW - PX - 0.7, IMG_Y + IMG_H - ty - 0.16, [[("Why it matters  ", {"bold": True, "color": TEAL, "size": 11.5})],
         [(takeaway, {"size": 11.5})]], anchor=MSO_ANCHOR.MIDDLE)
    notes(sl, title + "\n\n" + "\n".join(f"{i + 1}. {p[0] if isinstance(p, tuple) else p}: {p[1] if isinstance(p, tuple) else ''}" for i, p in enumerate(points))
          + "\n\n" + takeaway)
    return sl


def lifecycle_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "How a change moves through the portal", "Change lifecycle", n)
    steps = [("Request", "An operator (or a customer) asks: add, modify, remove, deploy, rotate the key…"),
             ("Approve", "Someone else approves it (four eyes). An emergency change needs a reason."),
             ("Window", "It runs now if a change window is open. Otherwise it is scheduled for the next window."),
             ("Run the job", "Snapshot, render, NAC and VyOS push, then verify. Progress and ETA are shown live."),
             ("Prove", "Every router registered, every host reachable, 73 tests, Nautobot == lab.conf."),
             ("Record", "The config diff per router, who asked, who approved and the outcome, kept per job.")]
    x, y, w, h = 0.45, 1.6, 2.05, 1.25
    for i, (t, b) in enumerate(steps):
        s = box(sl, x + i * (w + 0.06), y, w + (0.25 if i < 5 else 0), h, fill=TEAL if i % 2 == 0 else NAVY,
                shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        shape_text(s, t, size=15)
        text(sl, x + i * (w + 0.06) + 0.05, y + h + 0.2, w - 0.05, 1.6, b, size=12, color=INK)
    y2 = 4.55
    box(sl, 0.45, y2, 12.4, 2.25, fill=BG, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    text(sl, 0.7, y2 + 0.18, 6, 0.4, "What is never allowed", size=15, color=NAVY, bold=True)
    text(sl, 0.7, y2 + 0.6, 5.8, 1.6, ["• Approving your own request (four eyes)", "• A viewer or customer starting a job",
                                      "• Rotating the pre-shared key unless you are an admin", "• A secret reaching git, a job log or the screen"], size=12.5, spacing=4)
    text(sl, 6.9, y2 + 0.18, 6, 0.4, "What never needs approval", size=15, color=NAVY, bold=True)
    text(sl, 6.9, y2 + 0.6, 5.7, 1.6, ["• Adding a customer: it takes nothing away", "• Dry runs, tests and drift checks",
                                      "• Customer diagnostics from its own LAN host", "• Reading anything the role allows"], size=12.5, spacing=4)
    notes(sl, "The policy is editable in the portal: which job types need approval, change windows, emergency changes and expiry of stale requests.")


def results_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "Measured, not assumed", "Results in the lab", n)
    kpis = [("3 / 3", "hubs every spoke is registered with"), ("20 / 20", "LAN hosts reach every other host"),
            ("33 / 33", "application checks passing"), ("73 / 73", "Robot tests pass after each change"),
            ("0.6 s", "worst failover when a hub fails; 29 of 35 flows never noticed"), ("6.8 s", "worst failover for dual-homed flows when a provider fails"),
            ("54", "IKE SAs re-authenticated with the new pre-shared key"), ("5 m 28 s", "key rotation end to end, one router at a time, then verified")]
    x0, y0, cw, ch, gx, gy = 0.45, 1.35, 3.0, 1.75, 0.12, 0.16
    for i, (k, l) in enumerate(kpis):
        col, row = i % 4, i // 4
        x, y = x0 + col * (cw + gx), y0 + row * (ch + gy)
        box(sl, x, y, cw, ch, fill=BG, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        text(sl, x + 0.2, y + 0.15, cw - 0.4, 0.7, k, size=32, color=GREEN if i < 4 else TEAL, bold=True)
        text(sl, x + 0.2, y + 0.92, cw - 0.4, 0.8, l, size=12.5, color=INK)
    text(sl, 0.45, 5.25, 12.4, 1.5, [
        [("Provider failure: ", {"bold": True, "color": NAVY}), ("dual-homed customers (cust1, cust4) move to cloud 2 over the second provider. Single-homed customers lose that provider "
          "by design, and the portal shows it as 'cut off' rather than hiding it.", {})],
        [("Hub failure: ", {"bold": True, "color": NAVY}), ("customers keep two other hubs, live shortcuts need no hub, and routing converges when BGP gives up on the failed hub (hold 9 s).", {})]],
        size=12.5, spacing=6)
    notes(sl, "Figures from the lab on 27 September 2026: the health check, the failover measurements on the Resilience page, and the pre-shared key rotation job (CR-0021).")


def summary_slide(n):
    sl = prs.slides.add_slide(BLANK)
    box(sl, 0, 0, SW, SH, fill=NAVY)
    text(sl, 0.8, 0.8, 11, 0.8, "In short", size=34, color=WHITE, bold=True)
    pts = [("One source of truth", "lab.conf drives the configuration, Nautobot and the portal, and drift is detected against it."),
           ("Every change is a job", "It is planned, approved, run, verified and recorded with a before-and-after diff per router."),
           ("Resilience you can measure", "Simulate a failure and get the outage timeline for every flow."),
           ("Customers see their own service", "SLA, applications, requests and self-service diagnostics, and nothing else."),
           ("Secure by default", "Roles, four eyes, key rotation on demand, and secrets never shown.")]
    y = 1.9
    for h, b in pts:
        box(sl, 0.8, y + 0.08, 0.14, 0.5, fill=ORANGE)
        text(sl, 1.15, y, 11, 0.45, h, size=19, color=WHITE, bold=True)
        text(sl, 1.15, y + 0.42, 11, 0.45, b, size=14, color=RGBColor(0xCB, 0xD5, 0xE1))
        y += 0.98
    text(sl, 0.8, 6.85, 11, 0.4, "github.com/dcantor/c8000v-dmvpn-lab · portal on the lab host at :8094 · API at /docs", size=12, color=RGBColor(0x94, 0xA3, 0xB8))
    text(sl, 12.2, 7.08, 0.9, 0.3, str(n), size=10, color=RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.RIGHT)


def tasks_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "Every Provision task, at a glance", "Provision in depth", n)
    rows = [
        ("Add a customer", "A new site: router (C8000v or VyOS) and LAN host, provider link, tunnels to all three hubs, Nautobot. No hub changes.", "~15 min", "operator", "no"),
        ("Modify a customer", "Company, applications, preferred hub, site LAN, second provider, or router type (rebuilds the router).", "1–15 min", "operator", "yes"),
        ("Remove a customer", "Terraform forgets it, VMs deleted, out of lab.conf and Nautobot, provider port released.", "~3 min", "operator", "yes"),
        ("Deploy the model", "Re-render from lab.conf, terraform apply, push providers and VyOS, seed Nautobot, compare.", "~4 min", "operator", "yes"),
        ("Dry run", "terraform plan and Nautobot's rendering compared with lab.conf. Changes nothing.", "read only", "operator", "no"),
        ("Check for drift", "Every router's running configuration compared with the model.", "1–2 min", "operator", "no"),
        ("Fix drift", "Put every router back to the model, then check again.", "a few min", "operator", "yes"),
        ("Back up the lab", "One .tar.gz: lab.conf, customers, renders, Nautobot, running configs, Terraform state.", "~30 s", "operator", "no"),
        ("Restore from a backup", "Upload, see exactly what would change, then rebuild the lab to it.", "varies", "operator", "yes"),
        ("Rotate the pre-shared key", "A random 40-character key on every router; every IKE session re-made and verified.", "~7 min", "admin", "yes"),
        ("Run the tests", "The Robot Framework suites against the live lab, with the evidence kept.", "~10 min", "operator", "no"),
    ]
    cols = [(0.45, 2.45, "Task"), (2.95, 6.45, "What it does"), (9.45, 1.2, "Time"), (10.7, 1.1, "Who"), (11.85, 1.05, "Approval")]
    y = 1.3
    box(sl, 0.45, y, 12.45, 0.42, fill=NAVY)
    for x, w, t in cols:
        text(sl, x + 0.08, y + 0.07, w - 0.1, 0.3, t, size=12, color=WHITE, bold=True)
    y += 0.42
    rh = 0.445
    for i, r in enumerate(rows):
        box(sl, 0.45, y, 12.45, rh, fill=BG if i % 2 == 0 else WHITE)
        for j, ((x, w, _), v) in enumerate(zip(cols, r)):
            col = NAVY if j == 0 else (ORANGE if (j == 4 and v == "yes") or (j == 3 and v == "admin") else INK)
            text(sl, x + 0.08, y + 0.05, w - 0.12, rh - 0.06, v, size=12 if j == 0 else 10.5, color=col, bold=j == 0 or col == ORANGE)
        y += rh
    text(sl, 0.45, y + 0.12, 12.4, 0.4, "Approval follows the change policy (Jobs → Change policy…): four eyes, optional change windows, emergencies with a reason. "
         "Viewers and customers never start a job.", size=11, color=MUTED)
    notes(sl, "The eleven Provision tasks. Anything that could take something away needs a second person; adding and reading never do.")


def monitoring_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "How the lab is monitored", "Monitoring", n)
    text(sl, 0.45, 1.2, 12.3, 0.6, "The shared monitoring stack on the NMS covers this lab. IOS XE has no exporter, so the portal itself measures the "
         "C8000v routers every minute and publishes the results as Prometheus metrics.", size=14, color=MUTED)

    def node(x, y, w, h, title, sub, fill=WHITE, line=TEAL, tc=NAVY):
        box(sl, x, y, w, h, fill=fill, line=line, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(sl, x + 0.08, y + 0.08, w - 0.16, 0.32, title, size=12.5, color=tc, bold=True, align=PP_ALIGN.CENTER)
        text(sl, x + 0.08, y + 0.4, w - 0.16, h - 0.45, sub, size=10, color=MUTED if tc == NAVY else RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.CENTER)

    def arrow(x1, y1, x2, y2):
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = MUTED; c.line.width = Pt(1.4)
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))

    text(sl, 0.45, 2.0, 2.6, 0.3, "SOURCES", size=10, color=TEAL, bold=True)
    text(sl, 3.55, 2.0, 2.8, 0.3, "COLLECTION", size=10, color=TEAL, bold=True)
    text(sl, 6.85, 2.0, 2.6, 0.3, "STORES", size=10, color=TEAL, bold=True)
    text(sl, 10.2, 2.0, 2.6, 0.3, "VIEWS AND ALERTS", size=10, color=TEAL, bold=True)
    node(0.45, 2.35, 2.6, 0.95, "Catalyst 8000v", "hubs and customers")
    node(0.45, 3.5, 2.6, 0.95, "VyOS", "providers and VyOS customers")
    node(0.45, 4.65, 2.6, 0.95, "LAN hosts", "Alpine, one per site")
    node(3.55, 2.35, 2.8, 0.95, "The portal: /metrics", "polls every router each minute: NHRP, IPsec, BGP, CPU, SLA probes, app checks",
         fill=NAVY, line=NAVY, tc=WHITE)
    node(3.55, 3.5, 2.8, 0.95, "Exporters and Telegraf", "node-exporter, frr-exporter; Telegraf pushes")
    node(3.55, 4.65, 2.8, 0.95, "Syslog", "C8000v and VyOS → UDP 5514")
    node(6.85, 2.35, 2.6, 0.95, "Prometheus", "scrapes; targets from the portal's /api/sd")
    node(6.85, 3.5, 2.6, 0.95, "VictoriaMetrics", "long-term metrics")
    node(6.85, 4.65, 2.6, 0.95, "VictoriaLogs", "router syslog, searchable")
    node(10.2, 2.35, 2.7, 0.95, "Grafana", "C8000v DMVPN overview, node detail")
    node(10.2, 3.5, 2.7, 0.95, "Alerts", "Prometheus rules and vmalert on syslog")
    node(10.2, 4.65, 2.7, 0.95, "The portal", "SLA page and monthly PDF, from VictoriaMetrics")
    arrow(3.05, 2.82, 3.55, 2.82); arrow(3.05, 3.97, 3.55, 3.97); arrow(3.05, 4.1, 3.55, 5.0); arrow(3.05, 5.12, 3.55, 3.2)
    arrow(6.35, 2.82, 6.85, 2.82); arrow(6.35, 3.97, 6.85, 3.97); arrow(6.35, 5.12, 6.85, 5.12); arrow(8.15, 3.3, 8.15, 3.5)
    arrow(9.45, 2.82, 10.2, 2.82); arrow(9.45, 3.97, 10.2, 3.97); arrow(9.45, 5.0, 10.2, 3.2); arrow(9.45, 4.1, 10.2, 5.0)
    text(sl, 0.45, 5.95, 12.4, 1.0, [
        [("Alerts: ", {"bold": True, "color": NAVY}), ("DmvpnCustomerNotRegistered, DmvpnProviderSessionDown, DmvpnOverlayBgpDown, DmvpnRouterUnreachable, DmvpnLanHostDown, "
          "DmvpnRouterCpuHigh from metrics; BGP neighbour down, NHS down and IKE SA down from syslog.", {})],
        [("Maintenance: ", {"bold": True, "color": NAVY}), ("a job, a simulated failure or declared maintenance mutes the affected nodes' alerts and is left out of the SLA.", {})]],
        size=11.5, spacing=4)
    notes(sl, "Monitoring architecture. The portal is a Prometheus target like any exporter: lab_dmvpn_*, lab_bgp_*, lab_ipsec_*, lab_router_cpu_pct, lab_sla_*, lab_app_up.")



def nautobot_model_slide(n):
    sl = prs.slides.add_slide(BLANK)
    header(sl, "Where the customers and the VPN live in Nautobot", "Source of truth", n)
    text(sl, 0.45, 1.2, 12.3, 0.6, "lab.sh nautobot seed writes the whole lab into the shared Nautobot. The renderer can build every configuration from "
         "Nautobot alone, and a check proves the result is byte for byte the same as from lab.conf.", size=14, color=MUTED)

    def node(x, y, w, h, title, sub, fill=WHITE, line=TEAL, tc=NAVY):
        box(sl, x, y, w, h, fill=fill, line=line, lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(sl, x + 0.08, y + 0.07, w - 0.16, 0.32, title, size=12.5, color=tc, bold=True, align=PP_ALIGN.CENTER)
        text(sl, x + 0.08, y + 0.38, w - 0.16, h - 0.42, sub, size=10, color=MUTED if tc == NAVY else RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.CENTER)

    def line(x1, y1, x2, y2, label=None):
        c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = MUTED; c.line.width = Pt(1.3)
        if label:
            text(sl, (x1 + x2) / 2 - 0.9, (y1 + y2) / 2 - 0.22, 1.8, 0.25, label, size=9, color=MUTED, align=PP_ALIGN.CENTER)

    text(sl, 0.45, 1.95, 4, 0.3, "THE CUSTOMER", size=10, color=TEAL, bold=True)
    text(sl, 4.7, 1.95, 4, 0.3, "THE NETWORK", size=10, color=TEAL, bold=True)
    text(sl, 9.3, 1.95, 4, 0.3, "THE VPN SERVICE", size=10, color=TEAL, bold=True)
    node(0.45, 2.3, 3.6, 1.05, "Tenant", "the customer company; custom fields: account, address, contact, industry")
    node(0.45, 3.75, 3.6, 1.05, "Virtual servers", "one per application per hub, with its VIP; linked to tenants as 'Application subscriptions'")
    node(0.45, 5.2, 3.6, 0.95, "Locations", "c8000v-dmvpn-lab and regions c8d-east, -central, -west")
    node(4.7, 2.3, 4.0, 1.05, "Devices", "c8d- names; roles dmvpn-hub, dmvpn-spoke, wan-provider, host; tenant = customer", fill=NAVY, line=NAVY, tc=WHITE)
    node(4.7, 3.75, 4.0, 1.05, "Interfaces, IPs and cables", "Tunnel0 (cloud 1), Tunnel1 (cloud 2), Loopback10 (VIPs), access links; a cable per link")
    node(4.7, 5.2, 4.0, 0.95, "Prefixes with roles", "dmvpn-overlay 172.28 / 172.29, site-lan, wan-p2p, oob-management")
    node(9.3, 2.3, 3.6, 1.05, "Config context", "the DMVPN service: NHRP, network-ids, tunnel keys, IKEv2 / IPsec, providers. No secrets")
    node(9.3, 3.75, 3.6, 1.05, "BGP models", "AS 65100 / 65000 / 65010, a routing instance per router, a peering per session")
    node(9.3, 5.2, 3.6, 0.95, "Saved GraphQL query", "c8000v-dmvpn-lab-model: what the renderer reads")
    line(4.05, 2.82, 4.7, 2.82, "tenant"); line(4.05, 4.27, 4.7, 2.95); line(4.05, 5.67, 4.7, 3.1, "location")
    line(6.7, 3.35, 6.7, 3.75); line(6.7, 4.8, 6.7, 5.2)
    line(8.7, 2.82, 9.3, 2.82, "applies to"); line(8.7, 4.27, 9.3, 4.27, "per router")
    text(sl, 0.45, 6.35, 12.4, 0.6, [[("Deliberately not in Nautobot: ", {"bold": True, "color": NAVY}),
         ("the pre-shared key and the NHRP secret. They stay in secrets/ on the lab host.", {})]], size=11.5)
    notes(sl, "The Nautobot model. ./lab.sh nautobot render --check rebuilds every rendered file from the saved GraphQL query and compares it with lab.conf's.")


# ---------- build ----------
title_slide()
n = 2
problem_slide(n); n += 1
what_slide(n); n += 1
caps_slide(n); n += 1
TOUR = [
    ("01_login", "Access", "Everyone signs in, and the role decides what they see", [
        ("Sign in", "Staff roles are viewer, operator, approver and admin. A customer account sees only its own service. Sessions are signed cookies, passwords are PBKDF2 hashes.")],
     "Read-only users cannot start jobs, nobody approves their own change, and customers never see another customer."),
    ("02_cloud", "See", "The cloud: every router, read live", [
        ("KPIs", "Hubs, customers, registrations, shortcuts and BGP at a glance."),
        ("Every router", "NHRP registrations, IPsec, overlay BGP, the provider and the preferred hub, read from the routers themselves."),
        ("Both providers", "Access links and eBGP on mpls and mpls2."),
        ("Who you are", "The account, its role and the portal version, always visible.")],
     "One screen answers 'is every spoke registered with every hub?' with no CLI logins."),
    ("03_drift", "See", "Health and configuration drift", [
        ("Health", "Everything the model expects to be up is checked, and anything missing is named."),
        ("Drift", "Each router is compared with what lab.conf renders. 'Fix drift' puts it back as a controlled job.")],
     "Hand edits no longer go unnoticed: drift is visible per router, and fixing it is one approved job."),
    ("04_map", "See", "A topology map drawn from live state", [
        ("The map", "Hubs, providers, customers and LAN hosts. Tunnel colour shows health, and the preferred hub is drawn bold."),
        ("Layers", "Turn provider 2 and cloud 2, shortcuts, hosts and labels on or off."),
        ("Export", "SVG or PNG, light or dark, exactly as drawn."),
        ("Details", "Click a node to see the customer company, its applications (checked every minute) and the router's live state.")],
     "Everyone gets the same live picture, from operations to account managers and customers."),
    ("05_path", "See", "Path trace: watch the phase 3 shortcut form", [
        ("First via a hub, then direct", "The first packets go through a hub. NHRP then builds a spoke-to-spoke shortcut, drawn in purple."),
        ("Details and actions", "Trace, watch the shortcut form, ping every host, and view the config history per router."),
        ("Every trace, hop by hop", "The raw output is kept so you can copy or download it.")],
     "It shows that DMVPN phase 3 works end to end, and makes it easy to explain."),
    ("06_sla", "Prove", "SLA per customer, from VictoriaMetrics", [
        ("Against targets", "Availability, latency, loss and tunnel uptime, each compared with its SLA target."),
        ("Latency to each hub", "Probed every minute from the customer's own LAN host."),
        ("Window and report", "24 hours, 30 days or this month, with a monthly PDF report for the customer.")],
     "SLA reporting is automatic and based on measurements the customer can check."),
    ("07_resilience", "Prove", "Simulate failures on purpose", [
        ("Choose a failure", "A hub, a provider, or a customer circuit. The page says what should happen before you press anything."),
        ("Measure failover", "Every LAN host pings every other host five times a second while the fault is in."),
        ("Every measurement", "Flows unaffected, failed over or cut off, the worst failover and the time back to healthy.")],
     "Resilience becomes a number rather than a claim, for example 0.6 s worst case for a hub failure."),
    ("08_failover", "Prove", "One failover experiment, flow by flow", [
        ("The experiment", "What failed, for how long, and the headline figures."),
        ("Each flow", "The verdict and an outage timeline for every host pair, so you can see which paths moved and which were cut.")],
     "You can see exactly which customers a failure affects, and for how long."),
    ("09_provision", "Change", "Provision: every change is a job", [
        ("Tasks", "Eleven tasks: each shows what it does, how long it takes and who may run it. The next slides walk through them."),
        ("Pre-shared key", "Admins rotate the DMVPN key. The system picks a random 40-character key, and only its fingerprint is ever shown.")],
     "Operators pick what they want done, and the portal carries it out consistently on every router."),
    ("10_modify", "Change", "Modify a customer, with the plan shown first", [
        ("Router type", "Move a customer between C8000v and VyOS."),
        ("Preferred hub", "Pick the hub its traffic should prefer (local-preference 200)."),
        ("Second provider", "Dual-home the customer onto mpls2 and cloud 2."),
        ("The plan", "Exactly what will change on which routers, before anything runs.")],
     "No surprises: the impact is visible and reviewed before a change request is filed."),
    ("11_changes", "Control", "Change requests: four eyes and change windows", [
        ("Change requests", "Each request shows who asked, what for, who approved or rejected it, and its status, including requests from customers.")],
     "Disruptive work cannot start without a second person, and it runs only inside the agreed windows."),
    ("12_job", "Audit", "A job: who, what, and the proof", [
        ("Who", "Who started the job and who approved it."),
        ("Progress", "Live progress and the time it took."),
        ("What changed", "The before-and-after diff per router. Secrets are masked but their change is still visible."),
        ("Every step", "Each step's result, for example '54 IKE SAs on 8 routers, all made after the rotation'.")],
     "Every change leaves evidence: the steps, the approvals and the exact configuration difference."),
    ("13_tools", "Secure", "Lab tools and accounts", [
        ("Accounts and roles", "Admins add, disable and delete accounts and set passwords. Accounts still on their lab-default password are flagged."),
        ("Every tool", "Nautobot, the API, the lab hub and the repository, each with how to sign in.")],
     "Access is managed in one place, with visible warnings when the defaults have not been changed."),
    ("14_customer", "Customers", "The customer's own view", [
        ("Status", "Service status, incidents and planned maintenance for this customer."),
        ("Its own site", "The map, cut down to the customer's own site and the hubs it uses."),
        ("Its applications", "Its company details and live application status."),
        ("Its own pages", "Network, service levels, requests and diagnostics. Nothing else is reachable.")],
     "Customers can answer their own questions, which means fewer tickets and more trust."),
    ("15_requests", "Customers", "Customer self-service requests", [
        ("Ask for a change", "Applications, the preferred hub or a second provider, with a reason."),
        ("Track it", "Each request becomes a change request for the operators, and its progress is shown here.")],
     "Customers can request changes without email, and the request still goes through staff approval."),
    ("16_diag", "Customers", "Self-service diagnostics", [
        ("Tests from its LAN host", "The customer tests its applications, the hubs, or traces a path."),
        ("Every application, every hub", "Round-trip time, loss and HTTP status, run on demand.")],
     "The first-line diagnosis is done by the customer in seconds, with no ticket needed."),
    ("20_add", "Provision in depth", "Add a customer: everything allocated for you", [
        ("Router type", "Catalyst 8000v (NAC) or VyOS."),
        ("Preferred hub", "Or none: any hub."),
        ("Dual-homed", "A backup link into mpls2 and cloud 2."),
        ("Name and addresses", "The next free name, management IP and LAN."),
        ("Index, port, host", "Tunnel index, provider port, LAN host."),
        ("The company", "Pre-filled; goes to Nautobot as a tenant.")],
     "Nothing to look up or calculate: every address and port is allocated and checked free."),
    ("21_add_plan", "Provision in depth", "Add a customer: the plan, before anything runs", [
        ("Applications", "What the customer subscribes to; application checks start for them."),
        ("The plan", "Router, cloud, hub, WAN, LAN, what changes, and the capacity left afterwards. Here, a new VyOS router on both providers, with no hub changes."),
        ("Provision", "Starts the job, or files a change request if the policy asks for approval.")],
     "Onboarding a customer takes one form and about 15 minutes, with the whole impact visible up front."),
    ("22_remove", "Provision in depth", "Remove a customer, safely", [
        ("Which customer", "Its router type, region and addressing are shown."),
        ("What goes with it", "The company and its subscriptions are removed from customers.json and Nautobot."),
        ("The plan", "The router and host deleted, Terraform state, the provider port released, lab.conf and Nautobot updated."),
        ("Remove", "Destructive: it needs an approver under the default policy.")],
     "Decommissioning leaves nothing behind: no orphaned ports, config, Terraform state or Nautobot records."),
    ("23_restore", "Provision in depth", "Restore from a backup: see the difference first", [
        ("The backup", "A .tar.gz made by Back up the lab, uploaded here."),
        ("What is in it", "Version, date, host, files, every checksum verified, and its customers."),
        ("What would change", "Customers changed, added or removed, other settings, then NAC, pushes, Nautobot and verification."),
        ("Restore", "Only after approval; nothing changes before.")],
     "A known-good state is always one approved job away, and you know exactly what it will undo."),
    ("24_psk", "Provision in depth", "The DMVPN pre-shared key", [
        ("Key", "Where it is kept (secrets/ on the lab host) and its fingerprint. The key itself is never shown."),
        ("Set", "When it was last rotated, by whom, and the job."),
        ("Used by", "Every hub and customer router."),
        ("Rotate", "Admins only, with a confirmation, and through change control.")],
     "Rotating the key becomes routine: one button, a random key, and proof that every session uses it."),
    ("26_deploy", "Provision in depth", "Deploy the model: lab.conf onto every router", [
        ("The job", "Who started it, who approved it, 7 of 7 steps, 4 minutes."),
        ("The steps", "Snapshot, render, Terraform apply (NAC), the providers, Nautobot seeded, Nautobot's rendering compared with lab.conf, snapshot and compare. The log is kept below.")],
     "The same pipeline builds, fixes and restores the lab, so the result is always the model."),
    ("27_tests", "Provision in depth", "Run the tests: the lab proves itself", [
        ("The job", "Robot Framework suites against the live lab."),
        ("The result", "73 of 73 passed, in about 11 minutes."),
        ("The evidence", "The Robot report and log, and every router's configuration before and after.")],
     "Every change can end with a full test run, and the evidence is one click away."),
    ("29_maint", "Provision in depth", "Maintenance and the job list", [
        ("Maintenance", "Declared here or started automatically by disruptive jobs and simulated failures. Alerts are muted, customers told, and SLA leaves it out."),
        ("Declare", "What, and which nodes it affects."),
        ("Every job", "Status, a progress bar, duration and test results. A failed job resumes from the step that failed.")],
     "Planned work does not page anyone or count against the SLA, and customers see it coming."),
    ("28_policy", "Change control", "The change policy, edited in the portal", [
        ("The policy", "Which tasks need approval, four eyes, emergency approvals with a reason, expiry after 72 h, and change windows by day and time."),
        ("In one line", "The current policy is always shown above the requests.")],
     "The rules for change are visible and adjustable, not tribal knowledge."),
    ("30_history", "See", "A router's configuration history", [
        ("Router details", "Its live state: addresses, NHRP, IPsec, BGP, and its actions."),
        ("Every job that changed it", "Each job's diff for this router, newest first, with masked secrets.")],
     "When something changed on a router, you can see which job did it and exactly what it changed."),
    ("31_api", "Automate", "Everything is an API", [
        ("REST API", "Every screen in the portal is built on the same API, documented with Swagger."),
        ("By area", "Monitoring, state, customers, jobs, changes, SLA, and more.")],
     "Other tools and pipelines can drive the lab the way the portal does."),
    ("32_dark", "Access", "Light or dark, and roles everywhere", [
        ("Dark mode", "Every page, including the live map."),
        ("An operator", "It sees and changes the lab, but admin tasks such as key rotation and accounts are hidden.")],
     "The portal fits the user, and each role only sees the actions it may take."),
    ("40_graf_cloud", "Monitoring", "Grafana: the DMVPN cloud over the last 24 hours", [
        ("Portal jobs on the timeline", "Purple markers: every test run, failover and steering change the portal made."),
        ("Health", "The portal's verdict and firing alerts."),
        ("Registrations", "15 of 15: five customers, three hubs each."),
        ("Each customer, each hub", "Gaps line up with the failover experiments."),
        ("eBGP to the provider", "Per router, over time."),
        ("IPsec per router", "Sessions UP-ACTIVE.")],
     "What the portal shows now, Grafana shows over time, next to what the portal was doing."),
    ("41_graf_provider", "Monitoring", "Grafana: the overlay, the routers and the provider", [
        ("Overlay iBGP", "Established vs configured sessions per router."),
        ("C8000v CPU", "Measured by the portal over SSH: IOS XE has no exporter."),
        ("The provider", "node-exporter, frr-exporter and Telegraf: traffic on every access port."),
        ("Syslog events", "NHRP, IKE and IPsec messages per router, counted in VictoriaLogs.")],
     "Underlay and overlay on one page: a provider problem and its DMVPN symptoms line up."),
    ("42_graf_logs", "Monitoring", "Grafana: router syslog", [
        ("BGP Down from syslog", "Neighbour-down messages counted every 5 minutes per router."),
        ("Router syslog", "Every C8000v line, newest first, from VictoriaLogs."),
        ("The lab host", "CPU, memory, load and swap of the KVM host.")],
     "Logs and metrics side by side: no logging in to routers to read their buffers."),
    ("43_graf_lab", "Monitoring", "Grafana: the lab and the portal's jobs", [
        ("Host CPU", "All cores, per mode."),
        ("VMs running", "Every router, host and provider VM."),
        ("LAN hosts answering", "Reachability of each site's host."),
        ("The portal's jobs", "The last outcome of each job type; red was a failed run.")],
     "Even the automation is monitored: a failing job type shows up like any other fault."),
    ("44_graf_node", "Monitoring", "Grafana: the provider router in detail", [
        ("BGP peers", "8 of 8 established on mpls."),
        ("Access links", "Traffic on every customer and hub port."),
        ("CPU, memory, load", "From node-exporter and Telegraf.")],
     "Each VyOS router has full telemetry, down to each interface and BGP peer."),
    ("45_alerts", "Monitoring", "The alert rules", [
        ("Dmvpn* rules", "Six metric rules for this lab, all inactive: healthy. Three more rules watch syslog in vmalert. All are gated on the lab running and muted during maintenance.")],
     "Problems page someone before a customer calls, and planned work does not."),
    ("50_nb_devices", "Source of truth", "Nautobot: every device of the lab", [
        ("The lab's devices", "Routers, providers and LAN hosts, under their c8d- names in the shared Nautobot."),
        ("Tenant", "Each customer's router and host belong to its company."),
        ("Role", "dmvpn-hub, dmvpn-spoke, wan-provider, host."),
        ("Location", "Its region: c8d-east, -central or -west.")],
     "One inventory for every lab on the host; this lab's part is easy to find and to filter."),
    ("51_nb_tenant", "Source of truth", "Nautobot: a customer company as a tenant", [
        ("Tenant", "Prairie Grain Logistics, in the c8000v-dmvpn-lab customers group."),
        ("Custom fields", "Account, address, contact, email, industry, phone."),
        ("Subscribed applications", "A relationship to the virtual servers it uses, at every hub."),
        ("Its devices", "Its router and its LAN host.")],
     "The customer record the portal edits is the one Nautobot holds: no second customer list to keep in sync."),
    ("52_nb_apps", "Source of truth", "Nautobot: applications as virtual servers", [
        ("One per application per hub", "APP-1001 to APP-1010, each at the hubs that host it."),
        ("Its VIP", "An address on the hub's Loopback10, with port and protocol; the portal's application checks test these.")],
     "Applications and who subscribes to them are modelled, not just the routers."),
    ("53_nb_device", "Source of truth", "Nautobot: a hub", [
        ("The device", "Location, role dmvpn-hub, platform, management address, autonomous system."),
        ("Interfaces", "Its tunnels, loopbacks and access links."),
        ("Autonomous system", "AS 65100, from the BGP models; the Config Context tab shows the DMVPN service data that applies to it.")],
     "Everything needed to build this hub's configuration is on this page or one click away."),
    ("54_nb_ifaces", "Source of truth", "Nautobot: the VPN on the hub's interfaces", [
        ("Tunnel0", "Cloud 1: mGRE, phase 3, sourced from GigabitEthernet2."),
        ("Tunnel1", "Cloud 2, the backup: sourced from GigabitEthernet4 into mpls2."),
        ("Loopback10", "The hub LAN, carrying the application VIPs."),
        ("Access link", "Cabled to the provider's port.")],
     "The two DMVPN clouds are modelled as real interfaces with their sources, addresses and cables."),
    ("55_nb_prefixes", "Source of truth", "Nautobot: the addressing plan", [
        ("Cloud 1", "172.28.0.0/24, role dmvpn-overlay: network-id 1, tunnel key 100."),
        ("Cloud 2", "172.29.0.0/24: network-id 2, tunnel key 200."),
        ("Site LANs", "Each customer's LAN, role site-lan."),
        ("Roles", "site-lan, dmvpn-overlay, wan-p2p, oob-management.")],
     "Every prefix says what it is for, so the allocator and the renderer agree."),
    ("56_nb_tunnel_ips", "Source of truth", "Nautobot: tunnel addresses", [
        ("Tunnel0 addresses", "Each router's address in cloud 1, assigned to its Tunnel0: hubs .1 to .3, customers from their index.")],
     "NHRP registrations and BGP neighbours come straight from these addresses."),
    ("57_nb_peerings", "Source of truth", "Nautobot: every BGP session", [
        ("Peerings", "Hub↔hub, hub (route reflector) ↔ customer, and site ↔ provider, over both clouds and both providers.")],
     "The routing design is data: a missing or extra session shows up as a difference, not a surprise."),
    ("58_nb_bgp", "Source of truth", "Nautobot: BGP routing instances", [
        ("Routing instances", "One per router: its autonomous system (65100 sites, 65000 and 65010 providers) and router-id.")],
     "AS numbers and router-ids are allocated once and read from here."),
    ("59_nb_context", "Source of truth", "Nautobot: the DMVPN service as a config context", [
        ("The service", "AS, both clouds (overlay, network-id, tunnel key), NHRP hold times, hubs and spokes, the IKEv2 and IPsec proposals, the providers. No pre-shared key.")],
     "The VPN service's settings are one reviewed document, applied to the whole lab."),
    ("60_nb_graphql", "Source of truth", "Nautobot: the query the renderer reads", [
        ("Saved GraphQL query", "c8000v-dmvpn-lab-model: the config context, devices, interfaces, IPs and BGP in one call. The same renderer turns it into every configuration.")],
     "lab.sh nautobot render --check proves Nautobot and lab.conf produce identical configurations."),
    ("33_capacity", "Prove", "Capacity: hub load and room to grow", [
        ("Room to grow", "How many more C8000v or VyOS customers fit, and which limit runs out first: here, the provider's customer ports."),
        ("Each hub", "Spokes against the planning figure, IPsec against the platform maximum, CPU and DRAM against the router's own levels, WAN traffic against the licensed throughput."),
        ("Providers", "Customer ports used and free on each provider."),
        ("The lab host", "Memory, CPU, and vCPUs allocated against its cores, for every lab on the machine.")],
     "Growth is planned, not discovered: you see the next bottleneck before a hub or the lab fills up."),
]

SECTIONS = [
    ("divider", "A tour of the portal", "Numbered callouts on each screenshot match the explanations beside it."),
    "01_login", "02_cloud", "03_drift", "04_map", "05_path", "30_history", "06_sla", "07_resilience", "08_failover", "33_capacity",
    ("divider", "Provision in depth", "Every change is a job: planned, approved, run, verified and recorded."),
    "09_provision", tasks_slide, "20_add", "21_add_plan", "10_modify", "22_remove", "26_deploy", "23_restore",
    "24_psk", "12_job", "27_tests", "29_maint",
    ("divider", "Change control", "A second pair of eyes, change windows, and a record of everything."),
    "11_changes", "28_policy", lifecycle_slide,
    ("divider", "Customers, access and the API", "Each customer sees its own service; staff see what their role allows."),
    "14_customer", "15_requests", "16_diag", "13_tools", "31_api", "32_dark",
    ("divider", "Source of truth: Nautobot", "Where the customers, their applications and the VPN are modelled."),
    nautobot_model_slide, "50_nb_devices", "51_nb_tenant", "52_nb_apps", "53_nb_device", "54_nb_ifaces", "55_nb_prefixes",
    "56_nb_tunnel_ips", "57_nb_peerings", "58_nb_bgp", "59_nb_context", "60_nb_graphql",
    ("divider", "Monitoring", "Prometheus, VictoriaMetrics, VictoriaLogs and Grafana: the lab over time."),
    monitoring_slide, "40_graf_cloud", "41_graf_provider", "42_graf_logs", "43_graf_lab", "44_graf_node", "45_alerts",
]
BYNAME = {t[0]: t for t in TOUR}
EXTRA = {"12_job": [{"label": "steps", "box": [0.05, 0.085, 0.45, 0.375]}]}
for item in SECTIONS:
    if isinstance(item, tuple):
        divider(item[1], item[2], n)
    elif callable(item):
        item(n)
    else:
        name, kicker, title, pts, take = BYNAME[item]
        shot_slide(n, name, kicker, title, pts, take, extra_boxes=EXTRA.get(name))
    n += 1

results_slide(n); n += 1
summary_slide(n)

out = D / "c8000v-dmvpn-portal.pptx"
prs.save(out)
print(out, len(prs.slides), "slides")
