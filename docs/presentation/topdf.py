"""The deck as a PDF, drawn from the .pptx itself (no office suite needed): vector shapes and text, the screenshots as images.

It understands what build.py uses: rectangles, rounded rectangles, ovals, chevrons and pentagons, straight connectors
with arrowheads, pictures, and text boxes with mixed runs. Calibri is replaced by Liberation Sans, slightly condensed
to keep Calibri's line breaks."""
import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

D = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else D / "c8000v-dmvpn-portal.pptx"
OUT = SRC.with_suffix(".pdf")
FONTS = Path("/usr/share/fonts/truetype/liberation")
pdfmetrics.registerFont(TTFont("R", str(FONTS / "LiberationSans-Regular.ttf")))
pdfmetrics.registerFont(TTFont("B", str(FONTS / "LiberationSans-Bold.ttf")))
pdfmetrics.registerFont(TTFont("I", str(FONTS / "LiberationSans-Italic.ttf")))
SCALE = 0.93                      # Liberation Sans is wider than Calibri
PT = 1 / 12700                    # EMU → points

prs = Presentation(SRC)
PW, PH = prs.slide_width * PT, prs.slide_height * PT
c = canvas.Canvas(str(OUT), pagesize=(PW, PH))
c.setTitle("C8000v DMVPN Portal"); c.setAuthor("c8000v-dmvpn-lab"); c.setSubject("Problem, what the portal does, capabilities, annotated tour")


def rgb(colorfmt):
    try:
        return HexColor("#" + str(colorfmt.rgb))
    except Exception:                                           # noqa: BLE001
        return None


def geom(sh):
    x, y, w, h = sh.left * PT, sh.top * PT, sh.width * PT, sh.height * PT
    return x, PH - y - h, w, h                                   # reportlab: origin bottom-left


def fill_line(sh):
    f = rgb(sh.fill.fore_color) if sh.fill.type == MSO_FILL.SOLID else None
    ln = sh.line
    l = rgb(ln.color) if ln.fill.type == MSO_FILL.SOLID else None
    return f, l, (ln.width or 12700) * PT


def path_shape(sh, x, y, w, h):
    kind = sh.auto_shape_type
    p = c.beginPath()
    if kind in (MSO_SHAPE.CHEVRON, MSO_SHAPE.PENTAGON):
        d = min(w, h) * 0.5
        top, bot, mid = y + h, y, y + h / 2
        p.moveTo(x, top); p.lineTo(x + w - d, top); p.lineTo(x + w, mid); p.lineTo(x + w - d, bot); p.lineTo(x, bot)
        if kind == MSO_SHAPE.CHEVRON:
            p.lineTo(x + d, mid)
        p.close()
    return p


def draw_auto(sh):
    x, y, w, h = geom(sh)
    f, l, lw = fill_line(sh)
    if f is None and l is None:
        return
    c.saveState()
    if f is not None:
        c.setFillColor(f)
    if l is not None:
        c.setStrokeColor(l); c.setLineWidth(lw)
    kind = sh.auto_shape_type
    fillflag, strokeflag = int(f is not None), int(l is not None)
    if kind == MSO_SHAPE.OVAL:
        c.ellipse(x, y, x + w, y + h, stroke=strokeflag, fill=fillflag)
    elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        adj = sh.adjustments[0] if len(sh.adjustments) else 0.16667
        c.roundRect(x, y, w, h, min(w, h) * adj, stroke=strokeflag, fill=fillflag)
    elif kind in (MSO_SHAPE.CHEVRON, MSO_SHAPE.PENTAGON):
        c.drawPath(path_shape(sh, x, y, w, h), stroke=strokeflag, fill=fillflag)
    else:
        c.rect(x, y, w, h, stroke=strokeflag, fill=fillflag)
    c.restoreState()


def draw_connector(sh):
    x1, y1 = sh.begin_x * PT, PH - sh.begin_y * PT
    x2, y2 = sh.end_x * PT, PH - sh.end_y * PT
    col = rgb(sh.line.color) or HexColor("#888888")
    lw = (sh.line.width or 12700) * PT
    c.saveState(); c.setStrokeColor(col); c.setFillColor(col); c.setLineWidth(lw)
    ln = sh.line._get_or_add_ln()
    arrow = ln.find(qn("a:tailEnd")) is not None
    if arrow:                                                    # stop the line at the arrowhead's base
        import math
        L = math.hypot(x2 - x1, y2 - y1) or 1
        ux, uy = (x2 - x1) / L, (y2 - y1) / L
        a = max(6, lw * 4.5)
        bx, by = x2 - ux * a, y2 - uy * a
        c.line(x1, y1, bx, by)
        p = c.beginPath(); p.moveTo(x2, y2)
        p.lineTo(bx - uy * a * 0.45, by + ux * a * 0.45); p.lineTo(bx + uy * a * 0.45, by - ux * a * 0.45); p.close()
        c.drawPath(p, stroke=0, fill=1)
    else:
        c.line(x1, y1, x2, y2)
    c.restoreState()


def draw_picture(sh):
    x, y, w, h = geom(sh)
    c.drawImage(ImageReader(io.BytesIO(sh.image.blob)), x, y, w, h)
    ln = sh.line
    l, lw = (rgb(ln.color) if ln.fill.type == MSO_FILL.SOLID else None), (ln.width or 12700) * PT
    if l is not None:
        c.saveState(); c.setStrokeColor(l); c.setLineWidth(lw); c.rect(x, y, w, h, stroke=1, fill=0); c.restoreState()


def fontname(f):
    return "B" if f.bold else ("I" if f.italic else "R")


def draw_text(sh, is_shape):
    tf = sh.text_frame
    if not tf.text.strip():
        return
    x, y, w, h = geom(sh)
    ml, mr = (tf.margin_left or 91440) * PT, (tf.margin_right or 91440) * PT
    mt, mb = (tf.margin_top or 45720) * PT, (tf.margin_bottom or 45720) * PT
    maxw = w - ml - mr
    laid = []                                                    # [(lines, align, space_after)]; a line: (width, height, [(text, font, size, color)])
    for p in tf.paragraphs:
        toks = []
        for r in p.runs:
            f = r.font
            size = (f.size.pt if f.size else 18) * SCALE
            col = rgb(f.color) or HexColor("#000000")
            parts = r.text.split(" ")
            for i, wd in enumerate(parts):
                toks.append((wd, fontname(f), size, col, i < len(parts) - 1))
        if not toks:
            continue
        lines, cur, cw = [], [], 0.0
        lh = max(t[2] for t in toks) * 1.2
        for wd, fn, sz, col, sp in toks:
            ww = pdfmetrics.stringWidth(wd, fn, sz)
            spw = pdfmetrics.stringWidth(" ", fn, sz) if sp else 0
            if cur and cw + ww > maxw and wd:
                lines.append((cw, cur)); cur, cw = [], 0.0
            cur.append((wd, fn, sz, col, spw)); cw += ww + spw
        lines.append((cw, cur))
        laid.append(([(lw_, lh, items) for lw_, items in lines], p.alignment, (p.space_after.pt if p.space_after is not None else 0)))
    total = sum(sum(l[1] for l in ls) + sa for ls, _, sa in laid) - (laid[-1][2] if laid else 0)
    anchor = tf._bodyPr.get("anchor") or ("ctr" if is_shape else "t")
    top = PH - sh.top * PT - mt
    inner = h - mt - mb
    if anchor == "ctr":
        top -= (inner - total) / 2
    elif anchor == "b":
        top -= inner - total
    cy = top
    for lines, align, sa in laid:
        for lw_, lh, items in lines:
            base = cy - lh * 0.8
            tw = lw_ - (items[-1][4] if items else 0)
            if align == PP_ALIGN.CENTER:
                cx = x + ml + (maxw - tw) / 2
            elif align == PP_ALIGN.RIGHT:
                cx = x + ml + maxw - tw
            else:
                cx = x + ml
            for wd, fn, sz, col, spw in items:
                c.setFillColor(col); c.setFont(fn, sz); c.drawString(cx, base, wd)
                cx += pdfmetrics.stringWidth(wd, fn, sz) + spw
            cy -= lh
        cy -= sa


for sl in prs.slides:
    for sh in sl.shapes:
        t = sh.shape_type
        if t == MSO_SHAPE_TYPE.PICTURE:
            draw_picture(sh)
        elif sh.__class__.__name__ == "Connector":
            draw_connector(sh)
        elif t == MSO_SHAPE_TYPE.AUTO_SHAPE:
            draw_auto(sh)
            draw_text(sh, True)
        elif t == MSO_SHAPE_TYPE.TEXT_BOX:
            draw_text(sh, False)
    c.showPage()
c.save()
print(OUT, len(prs.slides), "pages")
